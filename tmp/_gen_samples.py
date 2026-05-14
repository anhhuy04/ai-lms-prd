"""Generate sample test files for AGENT_TEST_GUIDE TC-02..TC-11.

Uses xlsxwriter (more compatible with Dart `excel` package v4.0.6)
instead of openpyxl which produces files that throw
"Unexpected null value" on parse.
"""
import xlsxwriter
from docx import Document
from pathlib import Path

OUT = Path(__file__).parent

MCQ_HEADER = ["STT", "Câu hỏi", "A", "B", "C", "D", "Đáp án đúng", "Độ khó", "Tags"]


def make_mcq_excel(path: Path, rows: list[list]):
    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet("Trắc nghiệm")
    header_fmt = wb.add_format({
        "bold": True, "font_color": "white",
        "bg_color": "#1A6FAB", "text_wrap": True, "valign": "top",
    })
    for col_idx, header in enumerate(MCQ_HEADER):
        ws.write(0, col_idx, header, header_fmt)
    for row_idx, r in enumerate(rows, start=1):
        ws.write(row_idx, 0, row_idx)  # STT
        for col_idx, val in enumerate(r, start=1):
            ws.write(row_idx, col_idx, val)
    widths = [6, 60, 25, 25, 25, 25, 12, 8, 25]
    for col_idx, w in enumerate(widths):
        ws.set_column(col_idx, col_idx, w)
    wb.close()
    print(f"[ok] {path.name}")


# ── 1. mau_excel_mcq.xlsx — 5 câu MCQ có số liệu (mixed: math/physics) ──────
mcq_mixed = [
    [
        "Một xe đi quãng đường 60 km trong 2 giờ. Vận tốc trung bình của xe là bao nhiêu?",
        "20 km/h", "30 km/h", "40 km/h", "50 km/h", "B", 2, "vật lý, vận tốc"
    ],
    [
        "Tổng các góc trong một tam giác bằng bao nhiêu độ?",
        "90°", "180°", "270°", "360°", "B", 1, "hình học, tam giác"
    ],
    [
        "Một hình vuông có cạnh 5 cm. Diện tích hình vuông là bao nhiêu cm²?",
        "10", "15", "20", "25", "D", 2, "hình học, diện tích"
    ],
    [
        "Cho biểu thức 3x + 5 = 20. Giá trị của x là bao nhiêu?",
        "3", "5", "15", "25", "B", 2, "đại số, phương trình"
    ],
    [
        "Một vật có khối lượng 2 kg, gia tốc 3 m/s². Lực tác dụng lên vật là bao nhiêu Newton?",
        "1.5 N", "5 N", "6 N", "9 N", "C", 3, "vật lý, lực"
    ],
]
make_mcq_excel(OUT / "mau_excel_mcq.xlsx", mcq_mixed)

# ── 2. mau_excel_2mcq.xlsx — đúng 2 câu MCQ (test chip enable) ─────────────
make_mcq_excel(OUT / "mau_excel_2mcq.xlsx", mcq_mixed[:2])

# ── 3. mau_excel_1mcq.xlsx — chỉ 1 câu (test chip disable) ─────────────────
make_mcq_excel(OUT / "mau_excel_1mcq.xlsx", mcq_mixed[:1])

# ── 4. mau_excel_nonum.xlsx — MCQ không có số/đơn vị (test auto-downgrade) ─
mcq_nonum = [
    [
        "Thủ đô của nước Việt Nam là thành phố nào?",
        "Thành phố Hồ Chí Minh", "Hà Nội", "Đà Nẵng", "Huế", "B", 1, "địa lý, thủ đô"
    ],
    [
        "Tác giả của bài thơ 'Truyện Kiều' là ai?",
        "Nguyễn Du", "Hồ Xuân Hương", "Nguyễn Trãi", "Tố Hữu", "A", 1, "văn học, thơ"
    ],
    [
        "Loài động vật nào sau đây thuộc lớp Thú?",
        "Cá heo", "Cá voi sát thủ", "Cả hai đáp án trên", "Chỉ A đúng", "C", 2, "sinh học, động vật"
    ],
    [
        "Quá trình quang hợp diễn ra chủ yếu ở bộ phận nào của cây?",
        "Rễ", "Thân", "Lá", "Hoa", "C", 1, "sinh học, thực vật"
    ],
    [
        "Ai là người sáng lập ra Đảng Cộng sản Việt Nam?",
        "Hồ Chí Minh", "Trần Phú", "Lê Duẩn", "Võ Nguyên Giáp", "A", 1, "lịch sử"
    ],
]
make_mcq_excel(OUT / "mau_excel_nonum.xlsx", mcq_nonum)


# ── 5. mau_word_mcq.docx — Word MCQ với "Câu N:" format ────────────────────
def make_word_mcq(path: Path):
    doc = Document()
    doc.add_heading("Bộ câu hỏi trắc nghiệm mẫu — Toán & Vật lý", level=1)
    doc.add_paragraph("[ TRẮC NGHIỆM — Toán học ]")
    doc.add_paragraph("")

    questions = [
        {
            "q": "Một xe đi quãng đường 120 km trong 3 giờ. Vận tốc trung bình của xe là bao nhiêu?",
            "opts": ["20 km/h", "30 km/h", "40 km/h", "60 km/h"],
            "ans": "C",
        },
        {
            "q": "Tổng ba góc trong của một tam giác bằng bao nhiêu độ?",
            "opts": ["90°", "180°", "270°", "360°"],
            "ans": "B",
        },
        {
            "q": "Cho phương trình 2x + 4 = 10. Nghiệm x bằng bao nhiêu?",
            "opts": ["1", "2", "3", "4"],
            "ans": "C",
        },
        {
            "q": "Diện tích hình tròn bán kính r được tính theo công thức nào?",
            "opts": ["2πr", "πr²", "πd", "2r²"],
            "ans": "B",
        },
        {
            "q": "Một vật khối lượng 5 kg, gia tốc 2 m/s². Lực tác dụng F bằng bao nhiêu Newton?",
            "opts": ["2.5 N", "5 N", "7 N", "10 N"],
            "ans": "D",
        },
    ]

    for i, q in enumerate(questions, start=1):
        doc.add_paragraph(f"Câu {i}: {q['q']}")
        for letter, opt in zip("ABCD", q["opts"]):
            doc.add_paragraph(f"{letter}. {opt}")
        doc.add_paragraph(f"Đáp án: {q['ans']}")
        doc.add_paragraph("")

    doc.save(path)
    print(f"[ok] {path.name}")


make_word_mcq(OUT / "mau_word_mcq.docx")


# ── 6. kienthuc.docx — Bài giảng thuần, không có câu hỏi ──────────────────
def make_kienthuc(path: Path):
    doc = Document()
    doc.add_heading("Bài 1: Quang hợp ở thực vật", level=1)

    p1 = (
        "Quang hợp là quá trình mà cây xanh sử dụng năng lượng ánh sáng mặt trời "
        "để tổng hợp chất hữu cơ từ các chất vô cơ đơn giản như nước (H₂O) và "
        "khí cacbonic (CO₂). Đây là một trong những quá trình sinh học quan trọng "
        "nhất trên Trái Đất vì nó cung cấp nguồn năng lượng cho hầu hết các "
        "sinh vật sống và đồng thời giải phóng oxy vào khí quyển."
    )
    doc.add_paragraph(p1)

    doc.add_heading("1. Phương trình tổng quát", level=2)
    doc.add_paragraph(
        "Phương trình hóa học tổng quát của quá trình quang hợp được viết như sau: "
        "6 CO₂ + 6 H₂O → C₆H₁₂O₆ + 6 O₂. "
        "Trong đó, cacbonic và nước là nguyên liệu, glucose (đường) là sản phẩm chính, "
        "và oxy được giải phóng dưới dạng khí."
    )

    doc.add_heading("2. Bộ phận thực hiện quang hợp", level=2)
    doc.add_paragraph(
        "Quang hợp chủ yếu diễn ra trong lá cây, cụ thể là ở các tế bào chứa lục lạp. "
        "Lục lạp chứa chất diệp lục (chlorophyll) có màu xanh đặc trưng, có khả năng "
        "hấp thụ năng lượng ánh sáng — chủ yếu là ánh sáng đỏ và xanh lam, phản xạ "
        "ánh sáng xanh lục, đó là lý do vì sao lá cây có màu xanh."
    )

    doc.add_heading("3. Hai pha của quang hợp", level=2)
    doc.add_paragraph(
        "Quang hợp gồm hai pha chính: pha sáng và pha tối. Pha sáng cần ánh sáng "
        "trực tiếp, diễn ra trên màng tilacoit, có chức năng phân giải nước, "
        "tạo ATP, NADPH và giải phóng O₂. Pha tối (chu trình Calvin) không cần "
        "ánh sáng trực tiếp, diễn ra trong chất nền của lục lạp, sử dụng ATP và "
        "NADPH từ pha sáng để khử CO₂ thành glucose."
    )

    doc.add_heading("4. Các yếu tố ảnh hưởng đến quang hợp", level=2)
    doc.add_paragraph(
        "Có bốn yếu tố chính ảnh hưởng tới tốc độ quang hợp: cường độ ánh sáng, "
        "nồng độ khí CO₂, nhiệt độ và lượng nước. Khi cường độ ánh sáng tăng, "
        "tốc độ quang hợp tăng đến một mức bão hòa. Nhiệt độ tối ưu cho quang hợp "
        "thường trong khoảng 25–35°C đối với cây nhiệt đới. Thiếu nước sẽ làm "
        "khí khổng đóng lại, giảm hấp thu CO₂ và làm chậm quang hợp."
    )

    doc.add_heading("5. Ý nghĩa sinh học", level=2)
    doc.add_paragraph(
        "Quang hợp tạo ra chất hữu cơ cung cấp năng lượng cho toàn bộ chuỗi thức "
        "ăn trên cạn và đại dương. Lượng O₂ do quang hợp tạo ra duy trì sự sống "
        "cho động vật và con người. Đồng thời, quang hợp giúp giảm lượng CO₂ "
        "trong khí quyển, góp phần điều hòa khí hậu toàn cầu."
    )

    doc.save(path)
    print(f"[ok] {path.name}")


make_kienthuc(OUT / "kienthuc.docx")

print("\n✅ Done — 6 files created in tmp/")
